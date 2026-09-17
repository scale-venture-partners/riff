"""Turn a file (or raw text) into a Document of prose Blocks with source locations."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from riff.textutil import strip_inline_markdown, word_count

PROSE_KINDS = frozenset({"paragraph", "list_item", "table_cell", "notes"})
SUPPORTED = {".md", ".markdown", ".txt", ".text", ".html", ".htm", ".docx", ".pptx"}


@dataclass
class Block:
    text: str
    kind: str
    line: int
    col: int = 1
    label: str = ""
    level: int = 0
    bold_lead: bool = False
    # (line number, raw line) pairs so regex rules can report exact positions.
    raw_lines: list[tuple[int, str]] = field(default_factory=list)

    @property
    def is_prose(self) -> bool:
        return self.kind in PROSE_KINDS

    @property
    def word_count(self) -> int:
        return word_count(self.text)

    def location(self) -> str:
        return f"{self.label}:" if self.label else f"{self.line}:{self.col}"


@dataclass
class Document:
    path: str
    format: str
    blocks: list[Block]

    @property
    def prose(self) -> list[Block]:
        return [b for b in self.blocks if b.is_prose]

    @property
    def headings(self) -> list[Block]:
        return [b for b in self.blocks if b.kind in ("heading", "title")]

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks if b.text)

    @property
    def word_count(self) -> int:
        return sum(b.word_count for b in self.blocks)


def extract(path: str | Path) -> Document:
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f"{p}: unsupported file type {suffix!r} (supported: {', '.join(sorted(SUPPORTED))})")
    if suffix == ".docx":
        return extract_docx(p)
    if suffix == ".pptx":
        return extract_pptx(p)
    text = p.read_text(encoding="utf-8", errors="replace")
    if suffix in (".html", ".htm"):
        return extract_html(text, str(p))
    if suffix in (".md", ".markdown"):
        return extract_markdown(text, str(p))
    return extract_text(text, str(p))


# ---------------------------------------------------------------- plain text and markdown

_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_LIST_ITEM = re.compile(r"^(\s*)(?:[-*+]|\d+[.)])\s+(.*)$")
_FENCE = re.compile(r"^\s*(```|~~~)")
_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
_SETEXT = re.compile(r"^\s*(=+|-+)\s*$")
_BOLD_LEAD = re.compile(r"^\s*(?:\*\*|__)[^*_]+(?:\*\*|__)")


def extract_text(text: str, name: str = "<text>") -> Document:
    blocks: list[Block] = []
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        if not lines[i].strip():
            i += 1
            continue
        start = i
        while i < len(lines) and lines[i].strip():
            i += 1
        raw = [(n + 1, lines[n]) for n in range(start, i)]
        body = " ".join(ln.strip() for _, ln in raw)
        blocks.append(Block(text=body, kind="paragraph", line=start + 1, raw_lines=raw))
    return Document(path=name, format="txt", blocks=blocks)


def extract_markdown(text: str, name: str = "<markdown>") -> Document:
    lines = text.splitlines()
    blocks: list[Block] = []
    i = 0
    if lines and lines[0].strip() == "---":
        j = 1
        while j < len(lines) and lines[j].strip() not in ("---", "..."):
            j += 1
        i = j + 1
    in_fence = None
    while i < len(lines):
        line = lines[i]
        fence = _FENCE.match(line)
        if fence:
            if in_fence is None:
                in_fence = fence.group(1)
            elif fence.group(1) == in_fence:
                in_fence = None
            i += 1
            continue
        if in_fence or not line.strip() or _TABLE_ROW.match(line) or line.lstrip().startswith("<"):
            i += 1
            continue
        h = _HEADING.match(line)
        if h:
            blocks.append(
                Block(
                    text=strip_inline_markdown(h.group(2)),
                    kind="heading",
                    line=i + 1,
                    level=len(h.group(1)),
                    raw_lines=[(i + 1, line)],
                )
            )
            i += 1
            continue
        li = _LIST_ITEM.match(line)
        if li:
            start = i
            raw = [(i + 1, line)]
            i += 1
            while i < len(lines) and lines[i].strip() and not _LIST_ITEM.match(lines[i]) and lines[i][:1].isspace():
                raw.append((i + 1, lines[i]))
                i += 1
            body = " ".join([li.group(2).strip()] + [ln.strip() for _, ln in raw[1:]])
            blocks.append(
                Block(
                    text=strip_inline_markdown(body),
                    kind="list_item",
                    line=start + 1,
                    col=len(li.group(1)) + 1,
                    bold_lead=bool(_BOLD_LEAD.match(li.group(2))),
                    raw_lines=raw,
                )
            )
            continue
        start = i
        raw = []
        while i < len(lines) and lines[i].strip() and not _HEADING.match(lines[i]) and not _LIST_ITEM.match(lines[i]):
            if _FENCE.match(lines[i]) or _TABLE_ROW.match(lines[i]):
                break
            raw.append((i + 1, lines[i]))
            i += 1
        if len(raw) >= 2 and _SETEXT.match(raw[-1][1]):
            body = " ".join(ln.strip().lstrip("> ").strip() for _, ln in raw[:-1])
            blocks.append(
                Block(
                    text=strip_inline_markdown(body),
                    kind="heading",
                    line=start + 1,
                    level=1 if raw[-1][1].strip().startswith("=") else 2,
                    raw_lines=raw,
                )
            )
            continue
        if raw:
            body = " ".join(ln.strip().lstrip("> ").strip() for _, ln in raw)
            blocks.append(Block(text=strip_inline_markdown(body), kind="paragraph", line=start + 1, raw_lines=raw))
    return Document(path=name, format="md", blocks=blocks)


# ---------------------------------------------------------------- html

_HTML_BLOCKS = ("h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote", "td", "th", "dd", "dt",
                "figcaption", "title")


def extract_html(html: str, name: str = "<html>") -> Document:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "pre", "code"]):
        tag.decompose()
    blocks: list[Block] = []
    for el in soup.find_all(_HTML_BLOCKS):
        if el.find(_HTML_BLOCKS):
            continue
        text = el.get_text(" ", strip=True)
        text = re.sub(r"\s+", " ", text)
        if not text:
            continue
        line = getattr(el, "sourceline", None) or 0
        if el.name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            kind, level = "heading", int(el.name[1])
        elif el.name == "title":
            kind, level = "title", 0
        elif el.name == "li":
            kind, level = "list_item", 0
        elif el.name in ("td", "th"):
            kind, level = "table_cell", 0
        else:
            kind, level = "paragraph", 0
        first = next((c for c in el.children if getattr(c, "name", None) or str(c).strip()), None)
        bold_lead = getattr(first, "name", None) in ("strong", "b")
        blocks.append(
            Block(text=text, kind=kind, line=line, level=level, bold_lead=bold_lead, raw_lines=[(line, text)])
        )
    return Document(path=name, format="html", blocks=blocks)


# ---------------------------------------------------------------- docx and pptx


def extract_docx(path: Path) -> Document:
    import docx

    d = docx.Document(str(path))
    blocks: list[Block] = []
    for idx, para in enumerate(d.paragraphs, 1):
        text = re.sub(r"\s+", " ", para.text).strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        kind, level = "paragraph", 0
        if style.startswith("Heading"):
            kind = "heading"
            digits = re.findall(r"\d+", style)
            level = int(digits[0]) if digits else 1
        elif style == "Title":
            kind = "title"
        elif "List" in style:
            kind = "list_item"
        runs = [r for r in para.runs if r.text.strip()]
        bold_lead = bool(runs) and bool(runs[0].bold) and not all(bool(r.bold) for r in runs)
        blocks.append(
            Block(
                text=text,
                kind=kind,
                line=idx,
                label=f"paragraph {idx}",
                level=level,
                bold_lead=bold_lead,
                raw_lines=[(idx, text)],
            )
        )
    for t_idx, table in enumerate(d.tables, 1):
        for r_idx, row in enumerate(table.rows, 1):
            for c_idx, cell in enumerate(row.cells, 1):
                text = re.sub(r"\s+", " ", cell.text).strip()
                if text:
                    label = f"table {t_idx} r{r_idx}c{c_idx}"
                    blocks.append(Block(text=text, kind="table_cell", line=0, label=label, raw_lines=[(0, text)]))
    return Document(path=str(path), format="docx", blocks=blocks)


def extract_pptx(path: Path) -> Document:
    from pptx import Presentation

    prs = Presentation(str(path))
    blocks: list[Block] = []
    for s_no, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                for r_idx, row in enumerate(shape.table.rows, 1):
                    for c_idx, cell in enumerate(row.cells, 1):
                        text = re.sub(r"\s+", " ", cell.text).strip()
                        if text:
                            label = f"slide {s_no} table r{r_idx}c{c_idx}"
                            blocks.append(
                                Block(text=text, kind="table_cell", line=s_no, label=label, raw_lines=[(s_no, text)])
                            )
                continue
            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                continue
            is_title = title_shape is not None and shape.shape_id == title_shape.shape_id
            for para in shape.text_frame.paragraphs:
                text = re.sub(r"\s+", " ", "".join(r.text for r in para.runs)).strip()
                if not text:
                    continue
                kind = "title" if is_title else ("list_item" if para.level > 0 else "paragraph")
                runs = [r for r in para.runs if r.text.strip()]
                bold_lead = bool(runs) and bool(runs[0].font.bold) and not all(bool(r.font.bold) for r in runs)
                blocks.append(
                    Block(
                        text=text,
                        kind=kind,
                        line=s_no,
                        label=f"slide {s_no}",
                        level=para.level,
                        bold_lead=bold_lead,
                        raw_lines=[(s_no, text)],
                    )
                )
        if slide.has_notes_slide:
            notes = re.sub(r"\s+", " ", slide.notes_slide.notes_text_frame.text).strip()
            if notes:
                blocks.append(Block(text=notes, kind="notes", line=s_no, label=f"slide {s_no} notes",
                                    raw_lines=[(s_no, notes)]))
    return Document(path=str(path), format="pptx", blocks=blocks)
