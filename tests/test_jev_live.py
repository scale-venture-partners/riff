"""Live Jev test against the real API. Skipped unless TYPESAFE_API_KEY is set.

Per project convention, the model path is exercised against the real stack rather than mocked.
Run in CI only where the key is available; it is a couple of cheap calls.
"""

from __future__ import annotations

import os

import pytest

from riff.extract import extract_markdown
from riff.jev import run_jev
from riff.settings import Settings

pytestmark = pytest.mark.skipif(
    not os.environ.get("TYPESAFE_API_KEY"), reason="TYPESAFE_API_KEY not set; live Jev test skipped"
)


async def test_preamble_detected_live():
    doc = extract_markdown(
        "Before diving in, let me set up what follows and explain the shape of the argument."
    )
    settings = Settings(select=("JEV001",))
    findings, stats = await run_jev(doc, ["JEV001"], settings)
    assert stats["calls"] >= 1
    codes = {f.code for f in findings}
    assert "JEV001" in codes
    assert all(0.0 <= f.probability <= 1.0 for f in findings)


async def test_clean_prose_stays_quiet_live():
    doc = extract_markdown(
        "The service reads its configuration at startup and exits with an error when a variable is missing."
    )
    settings = Settings(select=("JEV001", "JEV101"))
    findings, _ = await run_jev(doc, ["JEV001", "JEV101"], settings)
    assert findings == [] or all(f.probability < 0.9 for f in findings)


async def test_strunk_white_rule_detected_live():
    doc = extract_markdown("The situation developed in an unsatisfactory manner over the relevant period.")
    settings = Settings(select=("JEV502",))
    findings, stats = await run_jev(doc, ["JEV502"], settings)
    assert stats["calls"] >= 1
    assert "JEV502" in {f.code for f in findings}


async def test_document_scope_rule_runs_once_live():
    doc = extract_markdown(
        "Hi there! I lead growth at Acme and noticed your team is scaling fast. "
        "We help companies streamline onboarding. About us: 500+ customers and a Series A. "
        "Would you be open to a quick chat next week?"
    )
    settings = Settings(select=("JEV010",))
    findings, stats = await run_jev(doc, ["JEV010"], settings)
    assert stats["calls"] == 1  # one call over the whole document, not per block
    assert "JEV010" in {f.code for f in findings}


async def test_formulaic_close_detected_live():
    doc = extract_markdown("Let me know if that works, and no worries if now is not the right time.")
    settings = Settings(select=("JEV111",))
    findings, _ = await run_jev(doc, ["JEV111"], settings)
    assert "JEV111" in {f.code for f in findings}


async def test_classifies_email_live():
    from riff.doctype import classify_document

    text = ("Hi Jordan, nice to meet you. I lead AI investing at the firm and wanted to reach out. "
            "Would you be up for a coffee sometime next week? Thanks, Alex")
    result = await classify_document(text, len(text.split()))
    assert result.source == "classified"
    assert result.type in {"email", "letter"}


async def test_form_specific_rule_fires_for_its_type_live():
    from riff.settings import Settings

    doc = extract_markdown("## v2.3.0\n\nVarious improvements and bug fixes. Minor changes throughout.")
    settings = Settings(select=("JEV660",), forced_type="release_notes")
    findings, _ = await run_jev(doc, ["JEV660"], settings)
    assert "JEV660" in {f.code for f in findings}


async def test_revision_artifact_detected_live():
    doc = extract_markdown("Note: this slide has been revised per your feedback to add Acme and Beta.")
    findings, _ = await run_jev(doc, ["JEV011"], Settings(select=("JEV011",)))
    assert "JEV011" in {f.code for f in findings}


async def test_dated_update_stamp_is_not_a_revision_artifact_live():
    doc = extract_markdown("Last updated 2026-09-01. Acme, Beta, and Gamma each doubled ARR over the past year.")
    findings, _ = await run_jev(doc, ["JEV011"], Settings(select=("JEV011",)))
    assert findings == []


# -- the narrative rules, at the levels above a paragraph ------------------------------------------

def _deck(tmp_path, slides):
    from pptx import Presentation
    from pptx.util import Inches, Pt

    from riff.extract import extract

    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for i, (text, size) in enumerate(((title, 36), (body, 18))):
            tf = slide.shapes.add_textbox(Inches(1), Inches(0.5 + 2 * i), Inches(8), Inches(1)).text_frame
            tf.text = text
            tf.paragraphs[0].runs[0].font.size = Pt(size)
    prs.save(str(tmp_path / "deck.pptx"))
    return extract(tmp_path / "deck.pptx")


async def test_a_slide_whose_body_ignores_its_title_is_flagged_live(tmp_path):
    doc = _deck(tmp_path, [
        ("Churn fell by half after the pricing change",
         "We hired twelve engineers in the third quarter and the new London office opens in March."),
        ("Churn fell by half after the pricing change",
         "Monthly churn went from 4.1% to 2.0% in the two quarters after annual plans replaced monthly ones."),
    ])
    findings, stats = await run_jev(doc, ["JEV702"], Settings(select=("JEV702",)))
    assert stats["units"] == {"section": 2}
    assert [f.label for f in findings] == ["slide 1"]


async def test_topic_label_titles_and_a_ghost_deck_are_flagged_live(tmp_path):
    labels = [("Market Context", "Spending on agent infrastructure grew through the year across our survey."),
              ("Portfolio Performance", "Revenue across the portfolio rose while burn fell in most companies."),
              ("Looking Ahead", "We will keep investing in the areas where we see the most momentum.")]
    doc = _deck(tmp_path, labels)
    findings, _ = await run_jev(doc, ["JEV701", "JEV711"], Settings(select=("JEV701", "JEV711")))
    codes = [f.code for f in findings]
    assert codes.count("JEV701") >= 2 and "JEV711" in codes


async def test_claim_titles_pass_the_ghost_deck_test_live(tmp_path):
    claims = [("Agent infrastructure spending doubled this year", "Survey of 400 buyers, spend up from $1.1M to $2.3M."),
              ("Our portfolio grew faster while burning less", "Median ARR growth 42%; burn multiple fell from 2.1x to 1.4x."),
              ("So we are putting two thirds of new capital into agents", "Three new investments this quarter.")]
    findings, _ = await run_jev(_deck(tmp_path, claims), ["JEV711"], Settings(select=("JEV711",)))
    assert findings == []
