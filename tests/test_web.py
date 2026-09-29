"""URL linting: fetch over real HTTP from a local server and keep only the article body."""

from __future__ import annotations

import io
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

from riff.cli import main
from riff.extract import extract_html
from riff.web import extract_url, fetch, is_url

BODY = (
    "The migration finished on Tuesday after three weeks of careful work by the platform team. "
    "Nothing else about the rollout needed a second pass."
)
PAGE = f"""<html><head><title>Migration story | Insights | Acme</title></head><body>
<header><a href="/">Acme</a></header>
<nav><ul><li>Products</li><li>Pricing</li><li>Careers</li></ul></nav>
<main>
  <h1>How we finished the migration</h1>
  <div class="rich">
    <h2>The plan</h2>
    <p>{BODY}</p>
    <p>A second paragraph explains what changed for customers during the cutover window.</p>
  </div>
  <section class="related"><h3>Related posts</h3><ul><li>Another long post title here</li></ul>
    <p>Subscribe to our newsletter for updates.</p></section>
</main>
<footer><p>Copyright Acme’s parent company. All rights reserved worldwide.</p></footer>
<script>var x = "not prose at all";</script>
</body></html>"""


PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _pptx_bytes() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Portfolio overview (+2 companies added)"
    slide.placeholders[1].text = "Acme and Beta each doubled ARR over the past year."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("A short paragraph of ordinary prose for the word document.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


class _Handler(BaseHTTPRequestHandler):
    def do_GET(self):
        if self.path == "/page":
            body, ctype = PAGE.encode(), "text/html; charset=utf-8"
        elif self.path == "/text":
            body, ctype = b"plain text", "text/plain"
        elif self.path == "/deck.pptx":
            body, ctype = _pptx_bytes(), PPTX_TYPE
        elif self.path == "/download/deck.pptx":
            body, ctype = _pptx_bytes(), "application/octet-stream"
        elif self.path == "/report.docx":
            body, ctype = _docx_bytes(), DOCX_TYPE
        elif self.path == "/fake.pptx":
            body, ctype = b"<html>login page</html>", "text/plain"
        elif self.path == "/junk.pptx":
            body, ctype = b"not a zip file", PPTX_TYPE
        else:
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args):
        pass


@pytest.fixture(scope="module")
def base_url():
    server = HTTPServer(("127.0.0.1", 0), _Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


def texts(doc):
    return [b.text for b in doc.blocks]


def test_is_url_only_matches_http_schemes():
    assert is_url("https://example.com/a") and is_url("HTTP://example.com")
    assert not is_url("notes.md") and not is_url("file:///etc/passwd") and not is_url("ftp://x/y")


def test_extract_url_keeps_title_and_body_drops_chrome(base_url):
    doc = extract_url(f"{base_url}/page")
    assert texts(doc) == [
        "How we finished the migration",
        "The plan",
        BODY,
        "A second paragraph explains what changed for customers during the cutover window.",
    ]
    assert [b.label for b in doc.blocks] == ["block 1", "block 2", "block 3", "block 4"]


def test_fetch_rejects_non_html_and_http_errors(base_url):
    with pytest.raises(ValueError, match="expected an HTML page or a .pptx/.docx file, got text/plain"):
        fetch(f"{base_url}/text")
    with pytest.raises(ValueError, match="HTTP 404"):
        fetch(f"{base_url}/missing")


def test_fetch_reports_unreachable_host():
    with pytest.raises(ValueError, match="could not fetch"):
        fetch("http://127.0.0.1:1/")


def test_main_content_prefers_largest_article():
    html = (
        "<body><article><p>Short teaser card for another story on the site.</p></article>"
        "<article><h1>Real story</h1><p>The real story has a much longer body than the teaser above it does.</p>"
        "</article></body>"
    )
    assert texts(extract_html(html, main_content=True)) == [
        "Real story",
        "The real story has a much longer body than the teaser above it does.",
    ]


def test_main_content_falls_back_to_body_without_article_or_main():
    html = (
        "<body><header><h1>Site banner heading</h1></header><h1>The post</h1>"
        "<p>Paragraph text that is long enough to count as real prose here.</p></body>"
    )
    assert texts(extract_html(html, main_content=True)) == [
        "The post",
        "Paragraph text that is long enough to count as real prose here.",
    ]


def test_main_content_keeps_h1_that_sits_outside_the_article_container():
    html = (
        "<body><h1>Outside title</h1><article><p>An article body paragraph with plenty of words in it.</p></article>"
        "</body>"
    )
    assert texts(extract_html(html, main_content=True))[0] == "Outside title"


def test_main_content_drops_role_based_chrome():
    html = (
        '<main><div role="navigation"><p>Skip to the main content of this page please.</p></div>'
        "<p>The one paragraph that is actually the content of this page.</p></main>"
    )
    assert texts(extract_html(html, main_content=True)) == ["The one paragraph that is actually the content of this page."]


def test_local_html_files_are_still_extracted_whole():
    doc = extract_html(PAGE)
    assert "Products" in texts(doc) and "Migration story | Insights | Acme" in texts(doc)
    assert all(b.label == "" for b in doc.blocks)


def test_cli_lints_url_and_flags_body_only(base_url, capsys):
    rc = main([f"{base_url}/page", "--no-jev", "--type", "blog_post", "--no-color"])
    out = capsys.readouterr().out
    assert rc == 0, out  # the curly quote lives only in the footer, which is not linted
    assert "All checks passed" in out


def test_cli_reports_fetch_failure_nonzero(capsys):
    rc = main(["http://127.0.0.1:1/", "--no-jev"])
    assert rc == 2
    assert "failed to lint" in capsys.readouterr().err


def test_extract_url_reads_pptx_by_content_type(base_url):
    doc = extract_url(f"{base_url}/deck.pptx")
    assert doc.path == f"{base_url}/deck.pptx"
    assert doc.format == "pptx"
    assert "Portfolio overview (+2 companies added)" in texts(doc)
    assert "Acme and Beta each doubled ARR over the past year." in texts(doc)


def test_extract_url_reads_pptx_served_as_octet_stream(base_url):
    assert "Portfolio overview (+2 companies added)" in texts(extract_url(f"{base_url}/download/deck.pptx"))


def test_extract_url_reads_docx(base_url):
    doc = extract_url(f"{base_url}/report.docx")
    assert doc.format == "docx"
    assert texts(doc) == ["A short paragraph of ordinary prose for the word document."]


def test_pptx_suffix_does_not_override_an_html_or_text_response(base_url):
    with pytest.raises(ValueError, match="got text/plain"):
        fetch(f"{base_url}/fake.pptx")


def test_corrupt_office_download_fails_loudly(base_url):
    with pytest.raises(ValueError, match="could not read as .pptx"):
        extract_url(f"{base_url}/junk.pptx")


def test_cli_lints_pptx_url_and_finds_revision_badge(base_url, capsys):
    rc = main([f"{base_url}/deck.pptx", "--no-jev", "--type", "memo", "--no-color", "--select", "RIF004"])
    out = capsys.readouterr().out
    assert rc == 1
    assert f"{base_url}/deck.pptx:slide 1: RIF004" in out
