"""The document hierarchy: how each format becomes document > sections > titles and paragraphs >
sentences, and how Jev rules at each level are planned onto it. Offline throughout; the live model
path is in test_jev_live."""

from __future__ import annotations

import pytest

from riff.doctype import from_format
from riff.engine import lint_document
from riff.extract import extract, extract_html, extract_markdown, extract_text
from riff.jev import DISCOURSE_MIN_WORDS, min_words, plan_units
from riff.rules import load_rules
from riff.rules.base import REGISTRY, Rule, register
from riff.settings import Settings

load_rules()

NESTED = """Lead-in before any heading.

# Top

Under top.

## Child A

A body.

### Grandchild

G body.

## Child B

B body.

# Second

S body.
"""


# -- building sections ---------------------------------------------------------------------------

def test_markdown_headings_nest_into_sections():
    doc = extract_markdown(NESTED)
    tree = [(s.level, s.title.text if s.title else None, [b.text for b in s.blocks]) for s in doc.walk_sections()]
    assert tree == [
        (0, None, ["Lead-in before any heading."]),
        (1, "Top", ["Under top."]),
        (2, "Child A", ["A body."]),
        (3, "Grandchild", ["G body."]),
        (2, "Child B", ["B body."]),
        (1, "Second", ["S body."]),
    ]
    top = doc.sections[1]
    assert [c.title.text for c in top.children] == ["Child A", "Child B"]
    assert top.body_text == "Under top.\n\nA body.\n\nG body.\n\nB body.", "a section's body includes its subsections"
    assert top.anchor is top.title


def test_the_outline_is_every_titled_section_in_order():
    assert [(lvl, t) for lvl, _, t in extract_markdown(NESTED).outline] == [
        (1, "Top"), (2, "Child A"), (3, "Grandchild"), (2, "Child B"), (1, "Second")]


def test_plain_text_is_one_untitled_section():
    doc = extract_text("One paragraph.\n\nAnother paragraph.")
    (only,) = doc.sections
    assert only.title is None and only.level == 0 and len(only.blocks) == 2
    assert only.anchor is only.blocks[0]


def test_a_document_title_names_the_document_not_a_section():
    doc = extract_html("<html><head><title>The Report</title></head><body><h1>Findings</h1><p>Body text.</p></body></html>")
    assert [s.title.text for s in doc.walk_sections()] == ["Findings"]


def test_docx_heading_styles_nest(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("A Report", style="Title")
    d.add_heading("Summary", level=1)
    d.add_paragraph("We recommend the plan.")
    d.add_heading("Detail", level=2)
    d.add_paragraph("The numbers behind it.")
    d.save(tmp_path / "r.docx")
    doc = extract(tmp_path / "r.docx")
    (summary,) = doc.sections
    assert summary.title.text == "Summary" and [c.title.text for c in summary.children] == ["Detail"]


@pytest.fixture
def make_deck(tmp_path):
    """A deck built from text boxes, as a slide library draws it: no title placeholders."""
    def make(slides, name="deck.pptx", notes=None, picture_on=()):
        import io

        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches, Pt

        prs = Presentation()
        for n, boxes in enumerate(slides, start=1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            for i, (text, size) in enumerate(boxes):
                tf = slide.shapes.add_textbox(Inches(1), Inches(0.5 + i), Inches(8), Inches(1)).text_frame
                tf.text = text
                for p in tf.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(size)
            if notes and n in notes:
                slide.notes_slide.notes_text_frame.text = notes[n]
            if n in picture_on:
                buf = io.BytesIO()
                Image.new("RGB", (40, 20), "red").save(buf, format="PNG")
                buf.seek(0)
                slide.shapes.add_picture(buf, Inches(1), Inches(4))
        path = tmp_path / name
        prs.save(str(path))
        return extract(path)
    return make


def test_each_slide_is_a_section_titled_by_its_largest_type(make_deck):
    doc = make_deck([
        [("FUND VIII", 11), ("Churn fell by half after the pricing change", 36), ("Monthly churn went from 4% to 2%.", 18)],
        [("$2.4B", 120), ("Aggregate portfolio ARR", 14)],
    ], notes={1: "Say the cohort number out loud."})
    one, two = doc.sections
    assert (one.label, one.title.text) == ("slide 1", "Churn fell by half after the pricing change")
    assert [b.text for b in one.blocks] == ["FUND VIII", "Monthly churn went from 4% to 2%.",
                                           "Say the cohort number out loud."], "notes belong to their slide"
    assert two.title is None, "a display figure is not a title"


def test_a_title_set_over_two_paragraphs_is_one_title(make_deck):
    doc = make_deck([[("One ask before Q2.\nConfirm timing", 40), ("Body copy under the title here.", 16)]])
    assert doc.sections[0].title.text == "One ask before Q2. Confirm timing"


def test_small_type_never_becomes_a_title(make_deck):
    doc = make_deck([[("A note in small print only", 12)]])
    assert doc.sections[0].title is None and doc.headings == []


def test_pictures_and_charts_are_counted_per_slide(make_deck):
    doc = make_deck([[("A title over a chart", 32)], [("A title over words", 32)]], picture_on=(1,))
    assert [s.visuals for s in doc.sections] == [1, 0]


def test_a_deck_is_a_presentation_by_format_and_a_forced_type_still_wins(make_deck):
    doc = make_deck([[("A title for one slide", 32)]])
    result = lint_document(doc, Settings(jev=False))
    assert (result.doc_type.type, result.doc_type.source) == ("presentation", "format")
    forced = lint_document(doc, Settings(jev=False, forced_type="report"))
    assert forced.doc_type.type == "report"
    assert from_format("md") is None


# -- planning Jev units at each level --------------------------------------------------------------

def _register(code, scope, view="text", fragments=False):
    if code not in REGISTRY:
        register(Rule(code=code, name=code.lower(), summary="s", category="Test", source="test", kind="jev",
                      scope=scope, view=view, question={"question": "q?"}, fragments=fragments))
    return code


PARA = _register("TSTP01", "block")
FRAG = _register("TSTP02", "block", fragments=True)
SENT = _register("TSTS01", "sentence")
TITLE = _register("TSTT01", "title")
SECT = _register("TSTC01", "section")
DOCT = _register("TSTD01", "document")
OUTL = _register("TSTD02", "document", view="outline")
STRU = _register("TSTD03", "document", view="structure")


def _units(doc, codes, **settings):
    return plan_units(doc, codes, Settings(**settings))


def test_the_word_floor_is_eight_for_prose_four_for_decks_or_the_setting(make_deck):
    md = extract_markdown("x")
    deck = make_deck([[("A title for one slide", 32)]])
    assert (min_words(md, Settings()), min_words(deck, Settings())) == (8, 4)
    assert min_words(md, Settings(jev_min_words=3)) == 3


def test_short_blocks_are_asked_only_the_rules_that_read_fragments():
    doc = extract_markdown("Five words in this block.\n\nThis paragraph has comfortably more than eight words in it.")
    units = _units(doc, [PARA, FRAG], jev_min_words=4)
    asked = {u.anchor.text.split()[0]: u.codes for u in units}
    assert asked == {"Five": [FRAG], "This": [PARA, FRAG]}
    assert DISCOURSE_MIN_WORDS == 8
    assert _units(doc, [PARA]) and all(u.anchor.text.startswith("This") for u in _units(doc, [PARA]))


def test_sentence_rules_get_one_unit_per_sentence():
    doc = extract_markdown("The first sentence is long enough to count here. Short one. "
                           "The third sentence is also long enough to count.")
    units = _units(doc, [SENT])
    assert [u.state for u in units] == ["The first sentence is long enough to count here.",
                                        "The third sentence is also long enough to count."]
    assert all(u.scope == "sentence" and u.snippet for u in units)


def test_title_units_carry_where_the_title_sits_and_what_it_heads(make_deck):
    doc = make_deck([[("What we will cover", 32), ("One. Two. Three.", 14)],
                     [("Churn fell by half after pricing", 32), ("Monthly churn went from 4% to 2% after annual plans.", 14)]])
    first, second = _units(doc, [TITLE])
    assert first.state == {"kind": "slide title", "title": "What we will cover", "position": "1 of 2",
                           "words_under_it": 3, "opening": "One. Two. Three."}
    assert second.state["position"] == "2 of 2"


def test_section_units_need_a_body_and_say_when_a_chart_is_the_evidence(make_deck):
    doc = make_deck([
        [("A divider with almost nothing under it", 32), ("Part two", 14)],
        [("Churn fell by half after the pricing change", 32),
         ("Monthly churn went from 4% to 2% in the two quarters after annual plans replaced monthly.", 14)],
    ], picture_on=(2,))
    (unit,) = _units(doc, [SECT])
    assert unit.state["title"] == "Churn fell by half after the pricing change"
    assert unit.state["position"] == "slide 2 of 2"
    assert unit.state["charts_or_images"].startswith("1 ")


def test_document_views_text_outline_and_structure():
    doc = extract_markdown("# Step 1: plan\n\nWrite the outline first.\n\n# Step 2: draft\n\nThen write each section."
                           "\n\n# Step 3: cut\n\nThen remove what doesn't carry the argument.")
    by_code = {tuple(u.codes): u for u in _units(doc, [DOCT, OUTL, STRU])}
    assert "Write the outline first." in by_code[(DOCT,)].state
    assert by_code[(OUTL,)].state["outline"] == "1. (section 1) Step 1: plan\n2. (section 2) Step 2: draft\n3. (section 3) Step 3: cut"
    assert "Step 2: draft -- Then write each section." in by_code[(STRU,)].state["sections"]


def test_an_outline_needs_three_titles_to_have_a_shape():
    doc = extract_markdown("# One\n\nBody.\n\n# Two\n\nBody.")
    assert _units(doc, [OUTL, STRU]) == []


def test_link_lists_are_not_prose_at_any_level():
    doc = extract_markdown("[a](http://x) · [b](http://y) · [c](http://z) · [d](http://w) and more links here")
    assert _units(doc, [PARA, SENT]) == []


# -- rules declare a level -------------------------------------------------------------------------

def test_a_rule_must_name_a_real_level_and_views_are_for_documents():
    with pytest.raises(ValueError, match="unknown scope"):
        register(Rule(code="TSTX01", name="x", summary="s", category="T", source="t", kind="jev", scope="chapter"))
    with pytest.raises(ValueError, match="only for document-scope"):
        register(Rule(code="TSTX02", name="x", summary="s", category="T", source="t", kind="jev", scope="section",
                      view="outline"))
    with pytest.raises(ValueError, match="unknown view"):
        register(Rule(code="TSTX03", name="x", summary="s", category="T", source="t", kind="jev",
                      scope="document", view="summary"))


def test_narrative_rules_sit_at_their_levels():
    levels = {c: (REGISTRY[c].scope, REGISTRY[c].view) for c in ("JEV701", "JEV702", "JEV703", "JEV711")}
    assert levels == {"JEV701": ("title", "text"), "JEV702": ("section", "text"),
                      "JEV703": ("section", "text"), "JEV711": ("document", "outline")}
    assert REGISTRY["JEV701"].applies_to == ("presentation",) and not REGISTRY["JEV703"].default
    assert REGISTRY["RIF005"].scope == "section" and REGISTRY["RIF005"].kind == "static"


def test_deck_conventions_are_not_letter_or_essay_tells():
    for code in ("JEV001", "JEV010", "JEV111", "JEV112"):
        assert "presentation" in REGISTRY[code].skip_for, code


def test_only_word_choice_rules_read_fragments():
    fragments = {c for c, r in REGISTRY.items() if r.fragments and not r.custom and not c.startswith("TST")}
    assert {"JEV301", "JEV302", "JEV304", "JEV502"} <= fragments
    assert not fragments & {"JEV001", "JEV103", "JEV207", "JEV008"}


# -- RIF005: numbered parts out of order -----------------------------------------------------------

def _steps(*titles):
    return extract_markdown("\n\n".join(f"# {t}\n\nBody of this section." for t in titles))


def test_a_part_after_a_later_one_is_flagged_once():
    out_of_order_parts = REGISTRY["RIF005"].check

    doc = _steps("Steps 1-4: capability", "Step 5: check", "Step 6: hooks", "Steps 5-6: check and gate", "Step 7: policy")
    (f,) = out_of_order_parts(doc, Settings())
    assert f.message == "Steps 5-6 comes after step 6 (9:1)" and f.snippet == "Steps 5-6: check and gate"


def test_ranges_overviews_and_separate_kinds_are_in_order():
    out_of_order_parts = REGISTRY["RIF005"].check

    assert out_of_order_parts(_steps("Steps 1-4", "Steps 1-2", "Steps 3-4", "Part 1", "Step 5"), Settings()) == []
    assert out_of_order_parts(_steps("Market overview", "Our plan"), Settings()) == []


def test_a_decks_eyebrow_carries_the_step_number(make_deck):
    out_of_order_parts = REGISTRY["RIF005"].check

    doc = make_deck([[("STEP 6: HOOKS", 11), ("A hook is code the model cannot skip", 32)],
                     [("STEPS 5-6: CHECK AND GATE", 11), ("Two new tools close the gaps", 32)]])
    (f,) = out_of_order_parts(doc, Settings())
    assert f.label == "slide 2" and "after step 6 (slide 1)" in f.message


def test_the_outline_flag_prints_the_hierarchy(tmp_path, capsys):
    from riff.cli import main

    p = tmp_path / "doc.md"
    p.write_text(NESTED)
    assert main([str(p), "--outline"]) == 0
    out = capsys.readouterr().out
    assert "(md: " in out and "  3:1: Top  [1 blocks, 8 words]" in out and "      11:1: Grandchild" in out
    assert main([str(tmp_path / "missing.md"), "--outline"]) == 2


def test_jev_min_words_flag_and_setting(tmp_path, capsys):
    from riff.cli import main
    from riff.settings import load_settings

    p = tmp_path / "doc.md"
    p.write_text("Fine.")
    assert main([str(p), "--no-jev", "--jev-min-words", "0"]) == 2
    (tmp_path / "riff.toml").write_text("jev-min-words = 3\n")
    assert load_settings(start=tmp_path).jev_min_words == 3
    (tmp_path / "riff.toml").write_text("jev-min-words = 0\n")
    with pytest.raises(ValueError, match="at least 1"):
        load_settings(start=tmp_path)


def test_list_rules_shows_each_rules_level(capsys):
    from riff.cli import main

    assert main(["--list-rules"]) == 0
    out = capsys.readouterr().out
    assert "JEV701  jev  title" in out and "JEV711  jev  document" in out and "RIF005       section" in out
    assert main(["--explain", "JEV711"]) == 0
    assert "scope: document (outline view)" in capsys.readouterr().out



# -- section dividers --------------------------------------------------------------------------------

def test_a_section_numbered_part_or_with_nothing_under_it_is_a_divider(make_deck):
    doc = make_deck([
        [("The eight-step pattern", 36), ("Part 01", 12), ("Each step is general.", 14)],
        [("The evidence", 36), ("Section 2 of the talk", 12)],
        [("Chapter iii closes the loop", 36), ("Chapter III", 12)],
        [("A title alone on its slide", 36)],
        [("Churn fell after the pricing change", 36), ("Monthly churn went from 4% to 2%.", 14)],
        [("Parts of the plan", 36), ("Parts arrive in two weeks from the supplier.", 14)],
    ])
    assert [s.is_divider for s in doc.sections] == [True, True, True, True, False, False]


def test_title_and_section_rules_skip_dividers(make_deck):
    doc = make_deck([
        [("The eight-step pattern", 36), ("Part 01", 12), ("The sections that follow show each step for decks.", 14)],
        [("Churn fell after the pricing change", 36), ("Monthly churn went from 4.1% to 2.0% after annual plans.", 14)],
    ])
    titles = [u.state["title"] for u in _units(doc, [TITLE])]
    sections = [u.state["title"] for u in _units(doc, [SECT])]
    assert titles == sections == ["Churn fell after the pricing change"]
    assert _units(doc, [TITLE])[0].state["position"] == "2 of 2", "positions still count the divider"
