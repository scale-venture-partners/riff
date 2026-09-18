"""Extraction: every supported format yields prose blocks with usable locations.

Fixtures are generated in a temp dir rather than read from committed sample files, so the
tests are self-contained (and runnable under mutation testing, which copies only source).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from riff.extract import extract, extract_html, extract_markdown, extract_text


@pytest.fixture
def make_docx():
    def _make(path: Path) -> Path:
        import docx

        d = docx.Document()
        d.add_heading("A Heading Here", level=1)
        p = d.add_paragraph()
        p.add_run("Lead").bold = True
        p.add_run(": the rest of the sentence is not bold.")
        d.add_paragraph("A second paragraph of ordinary prose for the body.")
        table = d.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "A table cell with words in it."
        d.save(str(path))
        return path

    return _make


@pytest.fixture
def make_pptx():
    def _make(path: Path) -> Path:
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "A Slide Title"
        slide.placeholders[1].text_frame.text = "A bullet of body text on the slide."
        slide.notes_slide.notes_text_frame.text = "Speaker notes with several words here."
        prs.save(str(path))
        return path

    return _make


def test_markdown_headings_and_lists():
    doc = extract_markdown("# Title Here\n\nA paragraph of prose.\n\n- one item\n- two item\n")
    assert doc.format == "md"
    assert [b.kind for b in doc.blocks] == ["heading", "paragraph", "list_item", "list_item"]
    assert doc.headings[0].text == "Title Here"
    assert doc.blocks[1].line == 3


def test_markdown_strips_inline_and_code_fence():
    doc = extract_markdown("Use **bold** and `code` and [a link](http://x.com).\n\n```\nnot prose\n```\n")
    assert doc.prose[0].text == "Use bold and code and a link."
    assert all("not prose" not in b.text for b in doc.blocks)


def test_markdown_front_matter_skipped():
    doc = extract_markdown("---\ntitle: x\n---\n\nReal content.\n")
    assert doc.prose[0].text == "Real content."


def test_hard_wrapped_phrase_spans_lines():
    from riff.rules.base import phrase_pattern, regex_matches

    doc = extract_text("please do this in order\nto finish today")
    block = doc.blocks[0]
    hits = list(regex_matches(block, phrase_pattern(["in order to"])))
    assert len(hits) == 1, "phrase wrapped across a newline should still match"


def test_html_blocks_and_bold_lead():
    doc = extract_html("<h1>Head</h1><p>Body text.</p><ul><li><strong>Key</strong>: value</li></ul>")
    kinds = {b.kind for b in doc.blocks}
    assert "heading" in kinds and "list_item" in kinds
    li = next(b for b in doc.blocks if b.kind == "list_item")
    assert li.bold_lead


def test_text_paragraphs_split_on_blank_lines():
    doc = extract_text("para one\nstill one\n\npara two\n")
    assert len(doc.blocks) == 2
    assert doc.blocks[1].line == 4


def test_setext_heading():
    doc = extract_markdown("Big Title\n=========\n\nBody paragraph here.\n")
    assert doc.headings[0].text == "Big Title"
    assert doc.headings[0].level == 1


def test_html_title_and_table_cells():
    doc = extract_html("<title>Page Title</title><table><tr><td>Cell one text</td></tr></table>")
    kinds = {b.kind for b in doc.blocks}
    assert "title" in kinds and "table_cell" in kinds


def test_extract_dispatches_by_extension(tmp_path):
    md = tmp_path / "a.md"
    md.write_text("# H\n\nSome prose here for the body.\n")
    assert extract(md).format == "md"
    txt = tmp_path / "a.txt"
    txt.write_text("Plain text paragraph with several words.\n")
    assert extract(txt).format == "txt"


def test_docx_roundtrip(tmp_path, make_docx):
    doc = extract(make_docx(tmp_path / "d.docx"))
    assert doc.format == "docx"
    assert {b.kind for b in doc.blocks} >= {"heading", "paragraph", "table_cell"}
    assert any(b.bold_lead for b in doc.blocks)
    assert any(b.label.startswith("paragraph") for b in doc.blocks)


def test_pptx_roundtrip(tmp_path, make_pptx):
    doc = extract(make_pptx(tmp_path / "p.pptx"))
    assert doc.format == "pptx"
    kinds = {b.kind for b in doc.blocks}
    assert "title" in kinds and "notes" in kinds
    assert any("slide 1" in b.label for b in doc.blocks)


def test_unsupported_extension(tmp_path):
    with pytest.raises(ValueError, match="unsupported"):
        extract(tmp_path / "nope.pdf")
